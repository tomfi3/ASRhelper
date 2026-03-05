"""Generate bias adjustment questionnaire Excel files for LAQM ASR reporting.

Creates one Excel workbook per monitoring station with:
- Sheet 1: Site & Study Information
- Sheet 2: Monthly Co-location Data (diffusion tube vs automatic analyser)
- Sheet 3: Bias Adjustment Calculation
- Sheet 4: Raw Hourly/Annual Statistics
"""

import logging
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

MONTHS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(bold=True, size=12)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
LIGHT_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
RESULT_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
WARN_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _style_header_row(ws, row, max_col):
    """Apply header styling to a row."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER


def _style_data_cell(ws, row, col, number_format=None):
    """Apply standard data cell styling."""
    cell = ws.cell(row=row, column=col)
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center")
    if number_format:
        cell.number_format = number_format
    return cell


def create_site_info_sheet(ws, site_code, site_info, year):
    """Create the Site & Study Information sheet."""
    ws.title = "Site & Study Info"
    ws.sheet_properties.tabColor = "4472C4"

    # Title
    ws.merge_cells("A1:D1")
    ws["A1"] = f"Bias Adjustment Questionnaire - {site_code}"
    ws["A1"].font = TITLE_FONT

    ws.merge_cells("A2:D2")
    ws["A2"] = f"Annual Status Report {year}"
    ws["A2"].font = SUBTITLE_FONT

    # Site details section
    row = 4
    ws.cell(row=row, column=1, value="SITE DETAILS").font = SUBTITLE_FONT
    row += 1

    fields = [
        ("Site Code", site_code),
        ("Site Name", site_info.get("name", "")),
        ("Site Type", site_info.get("site_type", "")),
        ("Location", site_info.get("location", "")),
        ("Borough", site_info.get("borough", "")),
        ("Monitoring Year", year),
    ]

    for label, value in fields:
        ws.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws.cell(row=row, column=1).fill = LIGHT_FILL
        ws.cell(row=row, column=1).border = THIN_BORDER
        ws.cell(row=row, column=2, value=value).border = THIN_BORDER
        row += 1

    # Diffusion tube details section
    row += 1
    ws.cell(row=row, column=1, value="DIFFUSION TUBE DETAILS").font = SUBTITLE_FONT
    row += 1

    tube_fields = [
        ("Tube Supplier/Laboratory", ""),
        ("Tube Preparation Method", ""),
        ("Tube Type", "Palmes-type"),
        ("Analysis Method", "Spectrophotometry"),
        ("Exposure Period", "Monthly (~4-5 weeks)"),
        ("Number of Tubes per Exposure", ""),
        ("Co-located with Automatic Analyser?", ""),
        ("Automatic Analyser Type", "Chemiluminescence"),
    ]

    for label, value in tube_fields:
        ws.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws.cell(row=row, column=1).fill = LIGHT_FILL
        ws.cell(row=row, column=1).border = THIN_BORDER
        cell = ws.cell(row=row, column=2, value=value)
        cell.border = THIN_BORDER
        if not value:
            cell.fill = WARN_FILL  # Highlight fields that need filling
        row += 1

    # Data source note
    row += 1
    ws.cell(row=row, column=1, value="DATA SOURCE").font = SUBTITLE_FONT
    row += 1
    ws.cell(row=row, column=1, value="Automatic analyser data sourced from London Air Quality Network API")
    ws.cell(row=row + 1, column=1, value="https://api.erg.ic.ac.uk/AirQuality/")

    # Column widths
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20


def create_monthly_data_sheet(ws, site_code, year, monthly_auto, monthly_dt):
    """Create the Monthly Co-location Data sheet.

    Args:
        ws: worksheet
        site_code: station code
        year: monitoring year
        monthly_auto: dict {month: {"mean": val, "data_capture": pct}} from automatic analyser
        monthly_dt: dict {month: value} from diffusion tubes
    """
    ws.title = "Monthly Data"
    ws.sheet_properties.tabColor = "70AD47"

    # Title
    ws.merge_cells("A1:H1")
    ws["A1"] = f"Monthly NO2 Concentrations (µg/m³) - {site_code} - {year}"
    ws["A1"].font = TITLE_FONT

    # Headers
    row = 3
    headers = [
        "Month",
        "Diffusion Tube\nNO2 (µg/m³)",
        "Automatic Analyser\nNO2 (µg/m³)",
        "Auto Data\nCapture (%)",
        "Tube/Auto\nRatio",
        "Difference\n(Tube - Auto)",
        "Tube\nExposure Start",
        "Tube\nExposure End",
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=header)
    _style_header_row(ws, row, len(headers))

    # Monthly data rows
    annual_tube_values = []
    annual_auto_values = []
    valid_month_count = 0

    for month_num in range(1, 13):
        row += 1
        month_name = MONTHS[month_num - 1]

        # Month name
        _style_data_cell(ws, row, 1)
        ws.cell(row=row, column=1, value=month_name)
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="left")

        # Diffusion tube value
        dt_val = monthly_dt.get(month_num)
        cell = _style_data_cell(ws, row, 2, "0.0")
        if dt_val is not None:
            cell.value = dt_val
            annual_tube_values.append(dt_val)

        # Automatic analyser mean
        auto_data = monthly_auto.get(month_num, {})
        auto_mean = auto_data.get("mean")
        cell = _style_data_cell(ws, row, 3, "0.0")
        if auto_mean is not None:
            cell.value = round(auto_mean, 1)
            annual_auto_values.append(auto_mean)

        # Data capture
        dc = auto_data.get("data_capture")
        cell = _style_data_cell(ws, row, 4, "0.0")
        if dc is not None:
            cell.value = dc
            if dc < 75:
                cell.fill = WARN_FILL

        # Ratio (tube/auto)
        cell = _style_data_cell(ws, row, 5, "0.00")
        if dt_val is not None and auto_mean is not None and auto_mean > 0:
            cell.value = round(dt_val / auto_mean, 2)

        # Difference
        cell = _style_data_cell(ws, row, 6, "0.0")
        if dt_val is not None and auto_mean is not None:
            cell.value = round(dt_val - auto_mean, 1)

        # Exposure dates (placeholder)
        _style_data_cell(ws, row, 7)
        _style_data_cell(ws, row, 8)

        if dt_val is not None and auto_mean is not None:
            valid_month_count += 1

    # Summary row
    row += 1
    ws.cell(row=row, column=1, value="Number of valid months").font = HEADER_FONT
    ws.cell(row=row, column=1).border = THIN_BORDER
    ws.cell(row=row, column=1).fill = RESULT_FILL
    cell = _style_data_cell(ws, row, 2)
    cell.value = valid_month_count
    cell.fill = RESULT_FILL

    # Annual means
    row += 1
    ws.cell(row=row, column=1, value="Annual Mean (raw)").font = HEADER_FONT
    ws.cell(row=row, column=1).border = THIN_BORDER
    ws.cell(row=row, column=1).fill = RESULT_FILL

    cell = _style_data_cell(ws, row, 2, "0.0")
    cell.fill = RESULT_FILL
    if annual_tube_values:
        cell.value = round(sum(annual_tube_values) / len(annual_tube_values), 1)

    cell = _style_data_cell(ws, row, 3, "0.0")
    cell.fill = RESULT_FILL
    if annual_auto_values:
        cell.value = round(sum(annual_auto_values) / len(annual_auto_values), 1)

    # Data capture summary
    row += 1
    ws.cell(row=row, column=1, value="Overall Data Capture (%)").font = HEADER_FONT
    ws.cell(row=row, column=1).border = THIN_BORDER
    ws.cell(row=row, column=1).fill = RESULT_FILL

    cell = _style_data_cell(ws, row, 2, "0.0")
    cell.fill = RESULT_FILL
    if annual_tube_values:
        cell.value = round((len(annual_tube_values) / 12) * 100, 1)

    cell = _style_data_cell(ws, row, 4, "0.0")
    cell.fill = RESULT_FILL
    all_dc = [monthly_auto.get(m, {}).get("data_capture", 0) for m in range(1, 13)]
    if any(all_dc):
        cell.value = round(sum(all_dc) / 12, 1)

    # Notes
    row += 2
    ws.cell(row=row, column=1, value="NOTES:").font = HEADER_FONT
    row += 1
    ws.cell(row=row, column=1, value="• Diffusion tube values are raw (unadjusted) monthly mean NO2 concentrations")
    row += 1
    ws.cell(row=row, column=1, value="• Automatic analyser values are monthly means calculated from hourly data")
    row += 1
    ws.cell(row=row, column=1, value="• Data capture < 75% highlighted in orange - annualisation may be required")
    row += 1
    ws.cell(row=row, column=1, value="• Ratio = Diffusion Tube / Automatic Analyser")

    # Column widths
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18


def create_bias_calculation_sheet(ws, site_code, year, monthly_auto, monthly_dt):
    """Create the Bias Adjustment Calculation sheet."""
    ws.title = "Bias Adjustment"
    ws.sheet_properties.tabColor = "ED7D31"

    # Title
    ws.merge_cells("A1:E1")
    ws["A1"] = f"Bias Adjustment Factor Calculation - {site_code} - {year}"
    ws["A1"].font = TITLE_FONT

    # Explanation
    row = 3
    ws.merge_cells(f"A{row}:E{row}")
    ws.cell(row=row, column=1, value=(
        "The bias adjustment factor accounts for the variable accuracy of diffusion tube "
        "NO2 measurements relative to the chemiluminescent reference method."
    ))

    # Method section
    row = 5
    ws.cell(row=row, column=1, value="CALCULATION METHOD").font = SUBTITLE_FONT
    row += 1
    ws.merge_cells(f"A{row}:E{row}")
    ws.cell(row=row, column=1, value=(
        "Local bias adjustment factor = Mean of (monthly auto analyser / monthly diffusion tube) ratios"
    ))

    # Data table
    row = 8
    headers = ["Month", "Tube (µg/m³)", "Auto (µg/m³)", "Auto/Tube Ratio", "Include?"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=h)
    _style_header_row(ws, row, len(headers))

    ratios = []
    for month_num in range(1, 13):
        row += 1
        _style_data_cell(ws, row, 1)
        ws.cell(row=row, column=1, value=MONTHS[month_num - 1])
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="left")

        dt_val = monthly_dt.get(month_num)
        auto_data = monthly_auto.get(month_num, {})
        auto_mean = auto_data.get("mean")

        cell = _style_data_cell(ws, row, 2, "0.0")
        if dt_val is not None:
            cell.value = dt_val

        cell = _style_data_cell(ws, row, 3, "0.0")
        if auto_mean is not None:
            cell.value = round(auto_mean, 1)

        cell = _style_data_cell(ws, row, 4, "0.000")
        include_cell = _style_data_cell(ws, row, 5)
        if dt_val is not None and auto_mean is not None and dt_val > 0:
            ratio = auto_mean / dt_val
            cell.value = round(ratio, 3)
            dc = auto_data.get("data_capture", 0)
            if dc >= 75:
                include_cell.value = "Yes"
                ratios.append(ratio)
            else:
                include_cell.value = "No (low DC)"
                include_cell.fill = WARN_FILL
        else:
            include_cell.value = "No data"
            include_cell.fill = WARN_FILL

    # Results section
    row += 2
    ws.cell(row=row, column=1, value="RESULTS").font = SUBTITLE_FONT
    row += 1

    result_fields = []

    if ratios:
        mean_ratio = sum(ratios) / len(ratios)
        bias_factor = round(mean_ratio, 2)

        # Calculate spread/precision
        if len(ratios) > 1:
            import math
            mean_r = mean_ratio
            variance = sum((r - mean_r) ** 2 for r in ratios) / (len(ratios) - 1)
            std_dev = math.sqrt(variance)
            # 95% CI using t-distribution approximation
            t_val = 2.0  # approximate for typical sample sizes
            ci = t_val * std_dev / math.sqrt(len(ratios))
        else:
            std_dev = 0
            ci = 0

        result_fields = [
            ("Number of valid paired months", len(ratios)),
            ("Mean Auto/Tube ratio (bias adjustment factor)", round(mean_ratio, 3)),
            ("Standard deviation of ratios", round(std_dev, 3) if std_dev else "N/A"),
            ("Approx. 95% CI of factor", f"±{round(ci, 3)}" if ci else "N/A"),
            ("", ""),
            ("TO APPLY BIAS ADJUSTMENT:", ""),
            ("Adjusted annual mean = Raw tube mean × bias adjustment factor", ""),
        ]

        # If we have annual means
        tube_values = [v for v in [monthly_dt.get(m) for m in range(1, 13)] if v is not None]
        if tube_values:
            raw_mean = sum(tube_values) / len(tube_values)
            adjusted_mean = raw_mean * mean_ratio
            result_fields.extend([
                ("", ""),
                ("Raw diffusion tube annual mean (µg/m³)", round(raw_mean, 1)),
                ("Bias adjustment factor", round(mean_ratio, 3)),
                ("Bias-adjusted annual mean (µg/m³)", round(adjusted_mean, 1)),
                ("", ""),
                ("Annual mean NO2 objective (µg/m³)", 40),
                ("Objective met?", "Yes" if adjusted_mean <= 40 else "No"),
            ])
    else:
        result_fields = [
            ("Number of valid paired months", 0),
            ("Insufficient data for local bias calculation", ""),
            ("Use national bias adjustment factor from LAQM spreadsheet", ""),
        ]

    for label, value in result_fields:
        ws.cell(row=row, column=1, value=label).font = HEADER_FONT if label else Font()
        ws.cell(row=row, column=1).border = THIN_BORDER if label else Border()
        if label:
            ws.cell(row=row, column=1).fill = RESULT_FILL
        cell = ws.cell(row=row, column=2, value=value)
        if label:
            cell.border = THIN_BORDER
            cell.fill = RESULT_FILL
            if isinstance(value, float):
                cell.number_format = "0.0" if value > 1 else "0.000"
        row += 1

    # National bias note
    row += 1
    ws.cell(row=row, column=1, value="NATIONAL BIAS ADJUSTMENT FACTOR").font = SUBTITLE_FONT
    row += 1
    ws.merge_cells(f"A{row}:E{row}")
    ws.cell(row=row, column=1, value=(
        "If a local bias factor cannot be calculated (fewer than 9 valid months), "
        "use the national combined bias adjustment factor from the LAQM national "
        "bias adjustment factors spreadsheet, matching your tube laboratory and "
        "preparation method. Download from: "
        "https://laqm.defra.gov.uk/air-quality/air-quality-assessment/national-bias/"
    ))
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True)

    # Column widths
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18


def create_annual_stats_sheet(ws, site_code, year, annual_stats, monthly_auto):
    """Create the Annual Statistics sheet with automatic analyser summary."""
    ws.title = "Annual Statistics"
    ws.sheet_properties.tabColor = "5B9BD5"

    # Title
    ws.merge_cells("A1:D1")
    ws["A1"] = f"Annual Air Quality Statistics - {site_code} - {year}"
    ws["A1"].font = TITLE_FONT

    # NO2 Statistics
    row = 3
    ws.cell(row=row, column=1, value="NO2 AUTOMATIC ANALYSER STATISTICS").font = SUBTITLE_FONT
    row += 1

    headers = ["Statistic", "Value", "Objective", "Met?"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=h)
    _style_header_row(ws, row, len(headers))

    annual_mean = annual_stats.get("annual_mean")
    data_capture = annual_stats.get("data_capture", 0)
    max_hourly = annual_stats.get("max_hourly")
    hours_above_200 = annual_stats.get("hours_above_200", 0)

    stats_rows = [
        ("Annual mean NO2 (µg/m³)", annual_mean, "40 µg/m³",
         "Yes" if annual_mean and annual_mean <= 40 else ("No" if annual_mean else "N/A")),
        ("Maximum hourly NO2 (µg/m³)", max_hourly, "200 µg/m³ (≤18 exceedances)",
         "N/A" if max_hourly is None else ("Yes" if max_hourly <= 200 else "Check")),
        ("Hours > 200 µg/m³", hours_above_200, "≤18 per year",
         "Yes" if hours_above_200 <= 18 else "No"),
        ("Data capture (%)", data_capture, "≥75% (min for valid assessment)", ""),
        ("Total valid hours", annual_stats.get("total_hours", 0), "", ""),
    ]

    for label, value, objective, met in stats_rows:
        row += 1
        ws.cell(row=row, column=1, value=label).font = HEADER_FONT
        ws.cell(row=row, column=1).border = THIN_BORDER

        cell = _style_data_cell(ws, row, 2)
        if value is not None:
            cell.value = value
            if isinstance(value, float):
                cell.number_format = "0.0"

        ws.cell(row=row, column=3, value=objective).border = THIN_BORDER
        met_cell = ws.cell(row=row, column=4, value=met)
        met_cell.border = THIN_BORDER
        if met == "No":
            met_cell.fill = WARN_FILL

    # Monthly breakdown table
    row += 2
    ws.cell(row=row, column=1, value="MONTHLY BREAKDOWN - AUTOMATIC ANALYSER NO2").font = SUBTITLE_FONT
    row += 1

    headers = ["Month", "Mean (µg/m³)", "Data Capture (%)", "Valid Hours"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=h)
    _style_header_row(ws, row, len(headers))

    for month_num in range(1, 13):
        row += 1
        _style_data_cell(ws, row, 1)
        ws.cell(row=row, column=1, value=MONTHS[month_num - 1])
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="left")

        auto_data = monthly_auto.get(month_num, {})

        cell = _style_data_cell(ws, row, 2, "0.0")
        mean_val = auto_data.get("mean")
        if mean_val is not None:
            cell.value = round(mean_val, 1)

        cell = _style_data_cell(ws, row, 3, "0.0")
        dc = auto_data.get("data_capture")
        if dc is not None:
            cell.value = dc
            if dc < 75:
                cell.fill = WARN_FILL

        cell = _style_data_cell(ws, row, 4)
        count = auto_data.get("count")
        if count is not None:
            cell.value = count

    # Column widths
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 15


def generate_bias_adjustment_workbook(
    site_code,
    site_info,
    year,
    monthly_auto,
    monthly_dt,
    annual_stats,
    output_dir="output",
):
    """Generate a complete bias adjustment questionnaire workbook for a site.

    Args:
        site_code: e.g. 'RI1'
        site_info: dict with name, site_type, location, borough
        year: monitoring year
        monthly_auto: dict {month: {"mean": val, "data_capture": pct, "count": n}}
        monthly_dt: dict {month: value} raw diffusion tube values
        annual_stats: dict from calculate_annual_stats()
        output_dir: directory for output files

    Returns:
        Path to created file
    """
    wb = Workbook()

    # Sheet 1: Site Info
    ws1 = wb.active
    create_site_info_sheet(ws1, site_code, site_info, year)

    # Sheet 2: Monthly Data
    ws2 = wb.create_sheet()
    create_monthly_data_sheet(ws2, site_code, year, monthly_auto, monthly_dt)

    # Sheet 3: Bias Adjustment Calculation
    ws3 = wb.create_sheet()
    create_bias_calculation_sheet(ws3, site_code, year, monthly_auto, monthly_dt)

    # Sheet 4: Annual Statistics
    ws4 = wb.create_sheet()
    create_annual_stats_sheet(ws4, site_code, year, annual_stats, monthly_auto)

    # Save
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    filename = output_path / f"BiasAdjustment_{site_code}_{year}.xlsx"
    wb.save(filename)
    logger.info(f"Created: {filename}")
    return filename
